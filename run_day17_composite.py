"""
SPRINT 3 - DAY 17
COMPOSITE QUALITY SCORE + SCREENER EXPORT

Run:
    python run_day17_composite.py
"""

from pathlib import Path
import sqlite3
import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter


# ============================================================
# PATHS
# ============================================================

ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / "db" / "nifty100.db"
OUTPUT_DIR = ROOT / "output"
OUTPUT_FILE = OUTPUT_DIR / "screener_output.xlsx"

OUTPUT_DIR.mkdir(exist_ok=True)


# ============================================================
# HELPERS
# ============================================================

def normalize_name(name):
    return (
        str(name)
        .strip()
        .lower()
        .replace(" ", "_")
        .replace("-", "_")
        .replace("/", "_")
        .replace("%", "pct")
        .replace("(", "")
        .replace(")", "")
    )


def clean_columns(df):
    df = df.copy()
    df.columns = [normalize_name(c) for c in df.columns]
    return df


def find_column(df, candidates):
    """
    Return first matching column from candidate list.
    """
    normalized = {normalize_name(c): c for c in df.columns}

    for candidate in candidates:
        c = normalize_name(candidate)

        if c in normalized:
            return normalized[c]

    return None


def numeric(df, column):
    if column is None or column not in df.columns:
        return pd.Series(np.nan, index=df.index)

    return pd.to_numeric(df[column], errors="coerce")


# ============================================================
# DATABASE
# ============================================================

def load_database():
    print("=" * 70)
    print("SPRINT 3 - DAY 17")
    print("COMPOSITE SCORE + EXPORT")
    print("=" * 70)

    if not DB_PATH.exists():
        raise FileNotFoundError(f"Database not found: {DB_PATH}")

    conn = sqlite3.connect(DB_PATH)

    tables = pd.read_sql_query(
        """
        SELECT name
        FROM sqlite_master
        WHERE type='table'
        ORDER BY name
        """,
        conn
    )

    print("\nDATABASE TABLES")
    print("-" * 70)

    print(tables["name"].tolist())

    ratio_tables = [
        x for x in tables["name"].tolist()
        if x.lower() == "financial_ratios"
    ]

    if not ratio_tables:
        conn.close()
        raise ValueError("financial_ratios table not found")

    ratios = pd.read_sql_query(
        "SELECT * FROM financial_ratios",
        conn
    )

    companies = pd.DataFrame()

    if "companies" in tables["name"].tolist():
        companies = pd.read_sql_query(
            "SELECT * FROM companies",
            conn
        )

    conn.close()

    ratios = clean_columns(ratios)

    if not companies.empty:
        companies = clean_columns(companies)

    print("\nFINANCIAL RATIO ROWS:", len(ratios))

    return ratios, companies


# ============================================================
# KEY NORMALIZATION
# ============================================================

def prepare_keys(df):
    df = df.copy()

    company_col = find_column(
        df,
        [
            "company_id",
            "id",
            "company",
            "company_code",
            "ticker",
            "symbol"
        ]
    )

    year_col = find_column(
        df,
        [
            "year",
            "financial_year",
            "fy",
            "fiscal_year"
        ]
    )

    if company_col is None:
        raise ValueError(
            "Could not identify company column. "
            f"Available columns: {list(df.columns)}"
        )

    if year_col is None:
        raise ValueError(
            "Could not identify year column. "
            f"Available columns: {list(df.columns)}"
        )

    df["company_id"] = df[company_col].astype(str).str.strip()

    df["year"] = (
        df[year_col]
        .astype(str)
        .str.extract(r"(\d{4})")[0]
    )

    df["year_num"] = pd.to_numeric(
        df["year"],
        errors="coerce"
    )

    return df


# ============================================================
# BUILD ANALYTICS DATASET
# ============================================================

def build_dataset(ratios, companies):
    print("\n" + "=" * 70)
    print("BUILDING SCREENER DATASET")
    print("=" * 70)

    ratios = prepare_keys(ratios)

    if not companies.empty:

        try:
            companies = prepare_keys(companies)

            company_cols = [
                c for c in companies.columns
                if c not in ratios.columns
                or c in ["company_id"]
            ]

            companies_small = companies[company_cols].copy()

            companies_small = (
                companies_small
                .drop_duplicates("company_id")
            )

            ratios = ratios.merge(
                companies_small,
                on="company_id",
                how="left",
                suffixes=("", "_company")
            )

        except Exception as e:
            print("Company merge skipped:", e)

    df = ratios.copy()

    # --------------------------------------------------------
    # METRIC MAPPING
    # --------------------------------------------------------

    roe_col = find_column(
        df,
        [
            "return_on_equity_pct",
            "roe_pct",
            "roe",
            "return_on_equity"
        ]
    )

    roce_col = find_column(
        df,
        [
            "return_on_capital_employed_pct",
            "roce_pct",
            "roce",
            "return_on_capital_employed"
        ]
    )

    npm_col = find_column(
        df,
        [
            "net_profit_margin_pct",
            "npm_pct",
            "net_profit_margin"
        ]
    )

    de_col = find_column(
        df,
        [
            "debt_to_equity",
            "de_ratio",
            "de"
        ]
    )

    fcf_col = find_column(
        df,
        [
            "free_cash_flow_cr",
            "free_cash_flow",
            "fcf_cr",
            "fcf"
        ]
    )

    rev5_col = find_column(
        df,
        [
            "revenue_cagr_5yr",
            "revenue_cagr_5_year",
            "revenue_cagr5",
            "sales_cagr_5yr",
            "sales_cagr"
        ]
    )

    pat5_col = find_column(
        df,
        [
            "pat_cagr_5yr",
            "net_profit_cagr_5yr",
            "profit_cagr_5yr"
        ]
    )

    eps5_col = find_column(
        df,
        [
            "eps_cagr_5yr",
            "eps_cagr"
        ]
    )

    icr_col = find_column(
        df,
        [
            "interest_coverage",
            "interest_coverage_ratio",
            "icr"
        ]
    )

    asset_turnover_col = find_column(
        df,
        [
            "asset_turnover"
        ]
    )

    payout_col = find_column(
        df,
        [
            "dividend_payout_ratio_pct",
            "dividend_payout_pct",
            "dividend_payout_ratio"
        ]
    )

    eps_col = find_column(
        df,
        [
            "earnings_per_share",
            "eps"
        ]
    )

    bvps_col = find_column(
        df,
        [
            "book_value_per_share",
            "bvps",
            "book_value"
        ]
    )

    market_price_col = find_column(
        df,
        [
            "market_price",
            "share_price",
            "stock_price",
            "close_price",
            "closing_price",
            "current_price",
            "price"
        ]
    )

    sales_col = find_column(
        df,
        [
            "sales",
            "revenue",
            "revenue_cr",
            "sales_cr",
            "net_sales"
        ]
    )

    net_profit_col = find_column(
        df,
        [
            "net_profit",
            "net_profit_cr",
            "pat",
            "profit_after_tax"
        ]
    )

    sector_col = find_column(
        df,
        [
            "broad_sector",
            "sector",
            "industry"
        ]
    )

    # --------------------------------------------------------
    # CREATE STANDARD COLUMNS
    # --------------------------------------------------------

    df["roe"] = numeric(df, roe_col)
    df["roce"] = numeric(df, roce_col)
    df["npm"] = numeric(df, npm_col)
    df["de"] = numeric(df, de_col)
    df["fcf"] = numeric(df, fcf_col)
    df["revenue_cagr_5yr"] = numeric(df, rev5_col)
    df["pat_cagr_5yr"] = numeric(df, pat5_col)
    df["eps_cagr_5yr"] = numeric(df, eps5_col)
    df["icr"] = numeric(df, icr_col)
    df["asset_turnover"] = numeric(df, asset_turnover_col)
    df["dividend_payout"] = numeric(df, payout_col)
    df["eps"] = numeric(df, eps_col)
    df["book_value"] = numeric(df, bvps_col)
    df["market_price"] = numeric(df, market_price_col)
    df["sales"] = numeric(df, sales_col)
    df["net_profit"] = numeric(df, net_profit_col)

    if sector_col:
        df["broad_sector"] = (
            df[sector_col]
            .astype(str)
            .str.strip()
        )
    else:
        df["broad_sector"] = "Unknown"

    # --------------------------------------------------------
    # DERIVE P/E
    # --------------------------------------------------------

    pe_existing = find_column(
        df,
        [
            "pe",
            "p_e",
            "pe_ratio",
            "price_to_earnings"
        ]
    )

    if pe_existing:
        df["pe"] = numeric(df, pe_existing)
    else:
        df["pe"] = np.where(
            (df["market_price"] > 0) &
            (df["eps"] > 0),
            df["market_price"] / df["eps"],
            np.nan
        )

    # --------------------------------------------------------
    # DERIVE P/B
    # --------------------------------------------------------

    pb_existing = find_column(
        df,
        [
            "pb",
            "p_b",
            "pb_ratio",
            "price_to_book"
        ]
    )

    if pb_existing:
        df["pb"] = numeric(df, pb_existing)
    else:
        df["pb"] = np.where(
            (df["market_price"] > 0) &
            (df["book_value"] > 0),
            df["market_price"] / df["book_value"],
            np.nan
        )

    # --------------------------------------------------------
    # DERIVE DIVIDEND YIELD
    # --------------------------------------------------------

    dividend_yield_existing = find_column(
        df,
        [
            "dividend_yield_pct",
            "dividend_yield",
            "yield_pct"
        ]
    )

    if dividend_yield_existing:
        df["dividend_yield"] = numeric(
            df,
            dividend_yield_existing
        )
    else:
        # Dividend yield = payout ratio * earnings yield
        #
        # payout is percentage, therefore:
        #
        # (payout / 100) * (EPS / Price) * 100
        #
        # = payout * EPS / Price

        df["dividend_yield"] = np.where(
            (df["dividend_payout"] >= 0) &
            (df["eps"] > 0) &
            (df["market_price"] > 0),

            df["dividend_payout"] *
            df["eps"] /
            df["market_price"],

            np.nan
        )

    # --------------------------------------------------------
    # DERIVE REVENUE CAGR 3 YEAR
    # --------------------------------------------------------

    rev3_existing = find_column(
        df,
        [
            "revenue_cagr_3yr",
            "revenue_cagr_3_year",
            "sales_cagr_3yr"
        ]
    )

    if rev3_existing:
        df["revenue_cagr_3yr"] = numeric(
            df,
            rev3_existing
        )
    else:
        df["revenue_cagr_3yr"] = np.nan

        # Try calculating from historical sales
        if sales_col:
            temp = df[
                ["company_id", "year_num", "sales"]
            ].copy()

            temp = temp.dropna(
                subset=["year_num", "sales"]
            )

            for company_id, group in temp.groupby(
                "company_id"
            ):

                group = group.sort_values("year_num")

                if len(group) < 4:
                    continue

                latest = group.iloc[-1]

                target_year = latest["year_num"] - 3

                previous = group.iloc[
                    (group["year_num"] - target_year)
                    .abs()
                    .argsort()
                ].iloc[0]

                start = previous["sales"]
                end = latest["sales"]

                if start > 0 and end > 0:
                    cagr = (
                        (end / start) ** (1 / 3) - 1
                    ) * 100

                    df.loc[
                        df["company_id"] == company_id,
                        "revenue_cagr_3yr"
                    ] = cagr

    # --------------------------------------------------------
    # FCF POSITIVE FLAG
    # --------------------------------------------------------

    df["fcf_positive"] = df["fcf"] > 0

    # --------------------------------------------------------
    # LATEST YEAR PER COMPANY
    # --------------------------------------------------------

    df = df.sort_values(
        ["company_id", "year_num"]
    )

    latest = (
        df.groupby("company_id", as_index=False)
        .tail(1)
        .copy()
    )

    print("\nLatest company rows:", len(latest))

    # --------------------------------------------------------
    # PRESET HELPER VALUES
    # --------------------------------------------------------

    # Debt-free
    latest["debt_free"] = latest["de"].fillna(
        np.inf
    ) == 0

    # Financials sector
    latest["is_financials"] = (
        latest["broad_sector"]
        .astype(str)
        .str.lower()
        .str.contains("financial")
    )

    # D/E decline
    de_history = df[
        ["company_id", "year_num", "de"]
    ].dropna()

    de_latest = (
        de_history
        .sort_values(["company_id", "year_num"])
        .groupby("company_id")
        .tail(2)
    )

    de_trend = {}

    for company_id, group in de_latest.groupby(
        "company_id"
    ):
        group = group.sort_values("year_num")

        if len(group) >= 2:
            old = group.iloc[-2]["de"]
            new = group.iloc[-1]["de"]

            de_trend[company_id] = new < old

    latest["de_declining"] = (
        latest["company_id"]
        .map(de_trend)
        .fillna(False)
    )

    # --------------------------------------------------------
    # DISPLAY / DIAGNOSTIC
    # --------------------------------------------------------

    print("\nDERIVED METRICS")
    print("-" * 70)

    diagnostic = [
        "roe",
        "roce",
        "npm",
        "de",
        "fcf",
        "revenue_cagr_3yr",
        "revenue_cagr_5yr",
        "pat_cagr_5yr",
        "eps_cagr_5yr",
        "pe",
        "pb",
        "dividend_yield",
        "dividend_payout",
        "sales",
        "icr"
    ]

    for col in diagnostic:
        if col in latest.columns:
            valid = latest[col].notna().sum()
            print(
                f"{col:<25} {valid:>4} valid / "
                f"{len(latest):>4}"
            )

    return df, latest


# ============================================================
# P10 / P90 NORMALIZATION
# ============================================================

def percentile_score(series):
    s = pd.to_numeric(
        series,
        errors="coerce"
    )

    if s.notna().sum() == 0:
        return pd.Series(
            50.0,
            index=series.index
        )

    p10 = s.quantile(0.10)
    p90 = s.quantile(0.90)

    if pd.isna(p10) or pd.isna(p90):
        return pd.Series(
            50.0,
            index=series.index
        )

    if p90 == p10:
        return pd.Series(
            50.0,
            index=series.index
        )

    clipped = s.clip(
        lower=p10,
        upper=p90
    )

    return (
        (clipped - p10) /
        (p90 - p10) *
        100
    ).fillna(0)


# ============================================================
# COMPOSITE SCORE
# ============================================================

def calculate_composite_score(df):

    print("\n" + "=" * 70)
    print("COMPOSITE QUALITY SCORE")
    print("=" * 70)

    df = df.copy()

    # --------------------------------------------------------
    # PROFITABILITY 35%
    # --------------------------------------------------------

    roe_score = percentile_score(df["roe"])
    roce_score = percentile_score(df["roce"])
    npm_score = percentile_score(df["npm"])

    profitability = (
        roe_score * 0.15 +
        roce_score * 0.10 +
        npm_score * 0.10
    )

    # --------------------------------------------------------
    # CASH QUALITY 30%
    # --------------------------------------------------------

    fcf_score = percentile_score(df["fcf"])

    # CFO/PAT
    cfo_col = find_column(
        df,
        [
            "cash_from_operations_cr",
            "operating_activity",
            "cash_from_operating_activity"
        ]
    )

    pat_col = find_column(
        df,
        [
            "net_profit",
            "net_profit_cr",
            "pat"
        ]
    )

    cfo = numeric(df, cfo_col)
    pat = numeric(df, pat_col)

    cfo_pat = np.where(
        pat != 0,
        cfo / pat,
        np.nan
    )

    df["cfo_pat_ratio"] = cfo_pat

    cfo_quality_score = percentile_score(
        df["cfo_pat_ratio"]
    )

    fcf_positive_score = np.where(
        df["fcf_positive"],
        100,
        0
    )

    cash_quality = (
        fcf_score * 0.15 +
        cfo_quality_score * 0.10 +
        fcf_positive_score * 0.05
    )

    # --------------------------------------------------------
    # GROWTH 20%
    # --------------------------------------------------------

    revenue_score = percentile_score(
        df["revenue_cagr_5yr"]
    )

    pat_growth_score = percentile_score(
        df["pat_cagr_5yr"]
    )

    growth = (
        revenue_score * 0.10 +
        pat_growth_score * 0.10
    )

    # --------------------------------------------------------
    # LEVERAGE 15%
    # --------------------------------------------------------

    # Lower D/E is better
    de_score = (
        100 -
        percentile_score(df["de"])
    )

    icr_score = percentile_score(
        df["icr"]
    )

    # Debt-free gets maximum ICR score
    icr_score = np.where(
        df["de"] == 0,
        100,
        icr_score
    )

    leverage = (
        de_score * 0.10 +
        icr_score * 0.05
    )

    # --------------------------------------------------------
    # FINAL
    # --------------------------------------------------------

    df["composite_quality_score"] = (
        profitability +
        cash_quality +
        growth +
        leverage
    )

    df["composite_quality_score"] = (
        df["composite_quality_score"]
        .clip(0, 100)
        .round(2)
    )

    print("✓ Composite score calculated")
    print(
        "Minimum:",
        round(
            df["composite_quality_score"].min(),
            2
        )
    )

    print(
        "Maximum:",
        round(
            df["composite_quality_score"].max(),
            2
        )
    )

    return df


# ============================================================
# PRESETS
# ============================================================

def run_presets(df):

    print("\n" + "=" * 70)
    print("RUNNING 6 PRESET SCREENERS")
    print("=" * 70)

    results = {}

    # --------------------------------------------------------
    # 1 QUALITY COMPOUNDER
    # --------------------------------------------------------

    mask = (
        (df["roe"] > 15) &
        (df["de"] < 1.0) &
        (df["fcf"] > 0) &
        (df["revenue_cagr_5yr"] > 10)
    )

    results["Quality Compounder"] = (
        df.loc[mask]
        .sort_values(
            "composite_quality_score",
            ascending=False
        )
        .copy()
    )

    # --------------------------------------------------------
    # 2 VALUE PICK
    # --------------------------------------------------------

    mask = (
        (df["pe"] > 0) &
        (df["pe"] < 20) &
        (df["pb"] > 0) &
        (df["pb"] < 3.0) &
        (df["de"] < 2.0) &
        (df["dividend_yield"] > 1)
    )

    results["Value Pick"] = (
        df.loc[mask]
        .sort_values(
            "composite_quality_score",
            ascending=False
        )
        .copy()
    )

    # --------------------------------------------------------
    # 3 GROWTH ACCELERATOR
    # --------------------------------------------------------

    mask = (
        (df["pat_cagr_5yr"] > 20) &
        (df["revenue_cagr_5yr"] > 15) &
        (df["de"] < 2.0)
    )

    results["Growth Accelerator"] = (
        df.loc[mask]
        .sort_values(
            "composite_quality_score",
            ascending=False
        )
        .copy()
    )

    # --------------------------------------------------------
    # 4 DIVIDEND CHAMPION
    # --------------------------------------------------------

    mask = (
        (df["dividend_yield"] > 2) &
        (df["dividend_payout"] < 80) &
        (df["fcf"] > 0)
    )

    results["Dividend Champion"] = (
        df.loc[mask]
        .sort_values(
            "composite_quality_score",
            ascending=False
        )
        .copy()
    )

    # --------------------------------------------------------
    # 5 DEBT-FREE BLUE CHIP
    # --------------------------------------------------------

    mask = (
        (df["de"] == 0) &
        (df["roe"] > 12) &
        (df["sales"] > 5000)
    )

    results["Debt-Free Blue Chip"] = (
        df.loc[mask]
        .sort_values(
            "composite_quality_score",
            ascending=False
        )
        .copy()
    )

    # --------------------------------------------------------
    # 6 TURNAROUND WATCH
    # --------------------------------------------------------

    mask = (
        (df["revenue_cagr_3yr"] > 10) &
        (df["fcf"] > 0) &
        (df["de_declining"])
    )

    results["Turnaround Watch"] = (
        df.loc[mask]
        .sort_values(
            "composite_quality_score",
            ascending=False
        )
        .copy()
    )

    # --------------------------------------------------------
    # PRINT
    # --------------------------------------------------------

    for name, result in results.items():
        print(
            f"\n{name}: "
            f"{len(result)} companies"
        )

    return results


# ============================================================
# EXPORT
# ============================================================

def export_excel(results):

    print("\n" + "=" * 70)
    print("GENERATING SCREENER EXCEL")
    print("=" * 70)

    # Keep the useful KPI columns.
    export_columns = [
        "company_id",
        "year",
        "broad_sector",

        "roe",
        "roce",
        "npm",

        "de",
        "icr",
        "fcf",

        "revenue_cagr_3yr",
        "revenue_cagr_5yr",
        "pat_cagr_5yr",
        "eps_cagr_5yr",

        "pe",
        "pb",
        "dividend_yield",
        "dividend_payout",

        "sales",
        "eps",
        "book_value",

        "asset_turnover",
        "cfo_pat_ratio",

        "composite_quality_score"
    ]

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl"
    ) as writer:

        for name, df in results.items():

            safe_name = name[:31]

            available = [
                c for c in export_columns
                if c in df.columns
            ]

            export_df = df[available].copy()

            export_df.to_excel(
                writer,
                sheet_name=safe_name,
                index=False
            )

    # --------------------------------------------------------
    # FORMATTING
    # --------------------------------------------------------

    wb = load_workbook(OUTPUT_FILE)

    green = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE"
    )

    red = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE"
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    for ws in wb.worksheets:

        # Header
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(
                bold=True
            )

        ws.freeze_panes = "A2"

        # Width
        for column_cells in ws.columns:

            max_length = 0
            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:
                try:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )
                except Exception:
                    pass

            ws.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                25
            )

        # Conditional formatting by useful metric
        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        # Positive metrics
        positive_columns = [
            "roe",
            "roce",
            "npm",
            "fcf",
            "revenue_cagr_5yr",
            "pat_cagr_5yr",
            "eps_cagr_5yr",
            "dividend_yield",
            "icr",
            "asset_turnover"
        ]

        for col_name in positive_columns:

            if col_name not in headers:
                continue

            col = headers[col_name]

            values = []

            for row in range(
                2,
                ws.max_row + 1
            ):
                value = ws.cell(
                    row=row,
                    column=col
                ).value

                try:
                    values.append(float(value))
                except Exception:
                    pass

            if not values:
                continue

            threshold = np.nanmedian(values)

            for row in range(
                2,
                ws.max_row + 1
            ):
                cell = ws.cell(
                    row=row,
                    column=col
                )

                try:
                    value = float(cell.value)

                    if value >= threshold:
                        cell.fill = green
                    else:
                        cell.fill = red

                except Exception:
                    pass

    wb.save(OUTPUT_FILE)

    print("\n✓ Excel created:")
    print(OUTPUT_FILE)


# ============================================================
# MAIN
# ============================================================

def main():

    ratios, companies = load_database()

    full_df, latest_df = build_dataset(
        ratios,
        companies
    )

    scored = calculate_composite_score(
        latest_df
    )

    results = run_presets(
        scored
    )

    export_excel(
        results
    )

    print("\n" + "=" * 70)
    print("DAY 17 SUMMARY")
    print("=" * 70)

    for name, result in results.items():
        print(
            f"{name:<25} "
            f"{len(result):>4} companies"
        )

    print("\nOutput:")
    print(OUTPUT_FILE)

    print("\n" + "=" * 70)
    print("DAY 17 COMPLETE")
    print("=" * 70)


if __name__ == "__main__":
    main()