"""
Sprint 3 - Day 17
Composite Quality Score + Screener Excel Export

Day 17 tasks:
1. Calculate composite quality score from 0-100.
2. Use P10/P90 winsorisation.
3. Calculate sector-relative scores.
4. Generate output/screener_output.xlsx.
5. Create one sheet per preset.
6. Sort by composite score descending.
7. Colour-code cells according to preset thresholds.
"""

from pathlib import Path
import sqlite3

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font


# ============================================================
# PROJECT PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

DATABASE_PATH = PROJECT_ROOT / "db" / "nifty100.db"
OUTPUT_DIR = PROJECT_ROOT / "output"
OUTPUT_FILE = OUTPUT_DIR / "screener_output.xlsx"


# ============================================================
# PRESETS
# ============================================================

PRESETS = {
    "Quality Compounder": {
        "roe": ("min", 15),
        "de": ("max", 1.0),
        "fcf": ("min", 0),
        "revenue_cagr_5yr": ("min", 10),
    },

    "Value Pick": {
        "pe": ("max", 20),
        "pb": ("max", 3.0),
        "de": ("max", 2.0),
        "dividend_yield": ("min", 1),
    },

    "Growth Accelerator": {
        "pat_cagr_5yr": ("min", 20),
        "revenue_cagr_5yr": ("min", 15),
        "de": ("max", 2.0),
    },

    "Dividend Champion": {
        "dividend_yield": ("min", 2),
        "dividend_payout": ("max", 80),
        "fcf": ("min", 0),
    },

    "Debt-Free Blue Chip": {
        "de": ("max", 0),
        "roe": ("min", 12),
        "sales": ("min", 5000),
    },

    "Turnaround Watch": {
        "revenue_cagr_3yr": ("min", 10),
        "fcf": ("min", 0),
    },
}


# ============================================================
# COMPOSITE ENGINE
# ============================================================

class CompositeEngine:

    def __init__(
        self,
        database_path=DATABASE_PATH
    ):

        self.database_path = Path(
            database_path
        )

        if not self.database_path.exists():

            raise FileNotFoundError(
                f"Database not found: "
                f"{self.database_path}"
            )

    # ========================================================
    # DATABASE
    # ========================================================

    def load_data(self):

        connection = sqlite3.connect(
            self.database_path
        )

        try:

            tables = pd.read_sql_query(
                """
                SELECT name
                FROM sqlite_master
                WHERE type='table'
                """,
                connection
            )["name"].tolist()

            if "financial_ratios" not in tables:

                raise ValueError(
                    "financial_ratios table not found."
                )

            df = pd.read_sql_query(
                "SELECT * FROM financial_ratios",
                connection
            )

            # ------------------------------------------------
            # Add company information if available
            # ------------------------------------------------

            if "companies" in tables:

                companies = pd.read_sql_query(
                    "SELECT * FROM companies",
                    connection
                )

                ratio_id = self.find_column(
                    df,
                    [
                        "company_id",
                        "id"
                    ]
                )

                company_id = self.find_column(
                    companies,
                    [
                        "company_id",
                        "id"
                    ]
                )

                if (
                    ratio_id is not None
                    and company_id is not None
                ):

                    company_columns = [
                        company_id
                    ]

                    sector = self.find_column(
                        companies,
                        [
                            "broad_sector",
                            "sector"
                        ]
                    )

                    company_name = self.find_column(
                        companies,
                        [
                            "company_name",
                            "name",
                            "company"
                        ]
                    )

                    if sector:
                        company_columns.append(sector)

                    if company_name:
                        company_columns.append(
                            company_name
                        )

                    company_info = (
                        companies[
                            company_columns
                        ]
                        .drop_duplicates(
                            subset=[company_id]
                        )
                    )

                    rename_map = {
                        company_id: ratio_id
                    }

                    if sector:
                        rename_map[
                            sector
                        ] = "broad_sector"

                    if company_name:
                        rename_map[
                            company_name
                        ] = "company_name"

                    company_info = company_info.rename(
                        columns=rename_map
                    )

                    # Don't overwrite existing columns
                    merge_columns = [
                        c for c in company_info.columns
                        if c == ratio_id
                        or c not in df.columns
                    ]

                    company_info = company_info[
                        merge_columns
                    ]

                    df = df.merge(
                        company_info,
                        on=ratio_id,
                        how="left"
                    )

            return df

        finally:

            connection.close()

    # ========================================================
    # FIND COLUMN
    # ========================================================

    @staticmethod
    def find_column(
        df,
        candidates
    ):

        lookup = {
            str(c).strip().lower(): c
            for c in df.columns
        }

        for candidate in candidates:

            key = str(candidate).strip().lower()

            if key in lookup:

                return lookup[key]

        return None

    # ========================================================
    # RESOLVE METRICS
    # ========================================================

    def resolve_metrics(
        self,
        df
    ):

        mappings = {

            "roe": [
                "return_on_equity_pct",
                "roe_pct",
                "roe",
                "roe_percentage"
            ],

            "roce": [
                "return_on_capital_employed_pct",
                "roce_pct",
                "roce",
                "roce_percentage"
            ],

            "npm": [
                "net_profit_margin_pct",
                "npm_pct",
                "npm"
            ],

            "fcf": [
                "free_cash_flow_cr",
                "free_cash_flow",
                "fcf"
            ],

            "fcf_cagr": [
                "fcf_cagr_5yr",
                "free_cash_flow_cagr_5yr"
            ],

            "cfo_pat": [
                "cfo_pat_ratio",
                "cfo_quality_ratio",
                "cfo_pat"
            ],

            "revenue_cagr": [
                "revenue_cagr_5yr"
            ],

            "pat_cagr": [
                "pat_cagr_5yr"
            ],

            "de": [
                "debt_to_equity",
                "de_ratio",
                "de"
            ],

            "icr": [
                "interest_coverage",
                "interest_coverage_ratio",
                "icr"
            ],
        }

        resolved = {}

        for name, candidates in mappings.items():

            column = self.find_column(
                df,
                candidates
            )

            if column is not None:

                resolved[name] = column

        return resolved

    # ========================================================
    # WINSORISED NORMALISATION
    # ========================================================

    @staticmethod
    def normalize_series(
        series
    ):

        values = pd.to_numeric(
            series,
            errors="coerce"
        )

        if values.notna().sum() == 0:

            return pd.Series(
                50.0,
                index=series.index
            )

        p10 = values.quantile(
            0.10
        )

        p90 = values.quantile(
            0.90
        )

        if pd.isna(p10) or pd.isna(p90):

            return pd.Series(
                50.0,
                index=series.index
            )

        if p90 == p10:

            return pd.Series(
                50.0,
                index=series.index
            )

        clipped = values.clip(
            lower=p10,
            upper=p90
        )

        return (
            (clipped - p10)
            / (p90 - p10)
            * 100
        )

    # ========================================================
    # SECTOR RELATIVE NORMALISATION
    # ========================================================

    def sector_normalize(
        self,
        df,
        column
    ):

        result = pd.Series(
            np.nan,
            index=df.index
        )

        sector_column = self.find_column(
            df,
            [
                "broad_sector",
                "sector"
            ]
        )

        if sector_column is None:

            return self.normalize_series(
                df[column]
            )

        for sector_name, indexes in df.groupby(
            sector_column
        ).groups.items():

            result.loc[indexes] = (
                self.normalize_series(
                    df.loc[indexes, column]
                )
            )

        return result

    # ========================================================
    # COMPOSITE SCORE
    # ========================================================

    def calculate_composite_score(
        self,
        df
    ):

        result = df.copy()

        metrics = self.resolve_metrics(
            result
        )

        print()
        print("=" * 70)
        print("COMPOSITE QUALITY SCORE")
        print("=" * 70)

        score_components = []

        # ----------------------------------------------------
        # Profitability: 35%
        # ----------------------------------------------------

        profitability = []

        for metric in [
            "roe",
            "roce",
            "npm"
        ]:

            if metric in metrics:

                normalized = self.sector_normalize(
                    result,
                    metrics[metric]
                )

                profitability.append(
                    normalized
                )

        if profitability:

            profitability_score = (
                profitability[0] * 0.15
                + (
                    profitability[1] * 0.10
                    if len(profitability) > 1
                    else 0
                )
                + (
                    profitability[2] * 0.10
                    if len(profitability) > 2
                    else 0
                )
            )

            score_components.append(
                profitability_score
            )

        # ----------------------------------------------------
        # Cash Quality: 30%
        # ----------------------------------------------------

        cash_components = []

        if "fcf_cagr" in metrics:

            cash_components.append(
                self.sector_normalize(
                    result,
                    metrics["fcf_cagr"]
                ) * 0.15
            )

        if "cfo_pat" in metrics:

            cash_components.append(
                self.sector_normalize(
                    result,
                    metrics["cfo_pat"]
                ) * 0.10
            )

        if "fcf" in metrics:

            fcf_values = pd.to_numeric(
                result[metrics["fcf"]],
                errors="coerce"
            )

            fcf_flag = (
                fcf_values > 0
            ).astype(float) * 100

            cash_components.append(
                fcf_flag * 0.05
            )

        if cash_components:

            cash_score = sum(
                cash_components
            )

            score_components.append(
                cash_score
            )

        # ----------------------------------------------------
        # Growth: 20%
        # ----------------------------------------------------

        growth_components = []

        if "revenue_cagr" in metrics:

            growth_components.append(
                self.sector_normalize(
                    result,
                    metrics["revenue_cagr"]
                ) * 0.10
            )

        if "pat_cagr" in metrics:

            growth_components.append(
                self.sector_normalize(
                    result,
                    metrics["pat_cagr"]
                ) * 0.10
            )

        if growth_components:

            growth_score = sum(
                growth_components
            )

            score_components.append(
                growth_score
            )

        # ----------------------------------------------------
        # Leverage: 15%
        # ----------------------------------------------------

        leverage_components = []

        if "de" in metrics:

            de_score = (
                100
                - self.sector_normalize(
                    result,
                    metrics["de"]
                )
            )

            leverage_components.append(
                de_score * 0.10
            )

        if "icr" in metrics:

            icr_score = self.sector_normalize(
                result,
                metrics["icr"]
            )

            leverage_components.append(
                icr_score * 0.05
            )

        if leverage_components:

            leverage_score = sum(
                leverage_components
            )

            score_components.append(
                leverage_score
            )

        # ----------------------------------------------------
        # Final score
        # ----------------------------------------------------

        if score_components:

            result[
                "composite_quality_score"
            ] = pd.concat(
                score_components,
                axis=1
            ).sum(
                axis=1,
                min_count=1
            )

        else:

            result[
                "composite_quality_score"
            ] = 0.0

        result[
            "composite_quality_score"
        ] = result[
            "composite_quality_score"
        ].clip(
            lower=0,
            upper=100
        )

        print(
            "✓ Composite score calculated"
        )

        print(
            f"Minimum: "
            f"{result['composite_quality_score'].min():.2f}"
        )

        print(
            f"Maximum: "
            f"{result['composite_quality_score'].max():.2f}"
        )

        return result

    # ========================================================
    # APPLY PRESET
    # ========================================================

    def apply_preset(
        self,
        df,
        preset_name
    ):

        result = df.copy()

        preset = PRESETS[preset_name]

        for metric_name, (
            operator,
            threshold
        ) in preset.items():

            column = self.get_preset_column(
                result,
                metric_name
            )

            if column is None:

                print(
                    f"⚠ {preset_name}: "
                    f"{metric_name} column not found"
                )

                continue

            values = pd.to_numeric(
                result[column],
                errors="coerce"
            )

            # ------------------------------------------------
            # Financials exception for D/E
            # ------------------------------------------------

            if metric_name == "de":

                sector_column = self.find_column(
                    result,
                    [
                        "broad_sector",
                        "sector"
                    ]
                )

                if sector_column:

                    sector = (
                        result[
                            sector_column
                        ]
                        .astype(str)
                        .str.lower()
                        .str.strip()
                    )

                    financials = (
                        sector == "financials"
                    )

                else:

                    financials = pd.Series(
                        False,
                        index=result.index
                    )

            else:

                financials = pd.Series(
                    False,
                    index=result.index
                )

            if operator == "min":

                mask = (
                    values >= threshold
                )

            else:

                mask = (
                    values <= threshold
                )

            if metric_name == "de":

                mask = (
                    mask
                    | financials
                )

            result = result.loc[
                mask.fillna(False)
            ].copy()

        return result

    # ========================================================
    # PRESET COLUMN RESOLVER
    # ========================================================

    def get_preset_column(
        self,
        df,
        metric
    ):

        mappings = {

            "roe": [
                "return_on_equity_pct",
                "roe_pct",
                "roe"
            ],

            "de": [
                "debt_to_equity",
                "de_ratio",
                "de"
            ],

            "fcf": [
                "free_cash_flow_cr",
                "free_cash_flow",
                "fcf"
            ],

            "revenue_cagr_5yr": [
                "revenue_cagr_5yr"
            ],

            "pat_cagr_5yr": [
                "pat_cagr_5yr"
            ],

            "pe": [
                "pe_ratio",
                "pe",
                "p_e"
            ],

            "pb": [
                "pb_ratio",
                "pb",
                "p_b"
            ],

            "dividend_yield": [
                "dividend_yield_pct",
                "dividend_yield"
            ],

            "dividend_payout": [
                "dividend_payout_ratio_pct",
                "dividend_payout_ratio",
                "dividend_payout"
            ],

            "sales": [
                "sales_cr",
                "sales",
                "revenue_cr",
                "revenue"
            ],

            "revenue_cagr_3yr": [
                "revenue_cagr_3yr"
            ]
        }

        return self.find_column(
            df,
            mappings.get(
                metric,
                [metric]
            )
        )

    # ========================================================
    # GENERATE EXCEL
    # ========================================================

    def generate_excel(
        self,
        all_results
    ):

        OUTPUT_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        print()
        print("=" * 70)
        print("GENERATING SCREENER EXCEL")
        print("=" * 70)

        with pd.ExcelWriter(
            OUTPUT_FILE,
            engine="openpyxl"
        ) as writer:

            for preset_name, df in all_results.items():

                if df is None:

                    continue

                sheet_name = (
                    preset_name[:31]
                )

                df = df.sort_values(
                    "composite_quality_score",
                    ascending=False
                )

                df.to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False
                )

        # ----------------------------------------------------
        # Formatting
        # ----------------------------------------------------

        self.format_excel(
            all_results
        )

        print(
            f"✓ Excel created:"
        )

        print(
            f"  {OUTPUT_FILE}"
        )

        return OUTPUT_FILE

    # ========================================================
    # FORMAT EXCEL
    # ========================================================

    def format_excel(
        self,
        all_results
    ):

        workbook = load_workbook(
            OUTPUT_FILE
        )

        green_fill = PatternFill(
            fill_type="solid",
            fgColor="C6EFCE"
        )

        red_fill = PatternFill(
            fill_type="solid",
            fgColor="FFC7CE"
        )

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        for preset_name, df in all_results.items():

            sheet_name = preset_name[:31]

            if sheet_name not in workbook.sheetnames:
                continue

            ws = workbook[
                sheet_name
            ]

            # ------------------------------------------------
            # Header
            # ------------------------------------------------

            for cell in ws[1]:

                cell.font = Font(
                    bold=True
                )

                cell.fill = header_fill

            # ------------------------------------------------
            # Find columns
            # ------------------------------------------------

            headers = {
                cell.value: cell.column
                for cell in ws[1]
            }

            preset = PRESETS[preset_name]

            for metric_name, (
                operator,
                threshold
            ) in preset.items():

                column = self.get_preset_column(
                    df,
                    metric_name
                )

                if column is None:
                    continue

                if column not in headers:
                    continue

                excel_column = headers[
                    column
                ]

                for row in range(
                    2,
                    ws.max_row + 1
                ):

                    value = ws.cell(
                        row=row,
                        column=excel_column
                    ).value

                    try:

                        value = float(value)

                    except (
                        TypeError,
                        ValueError
                    ):

                        continue

                    if operator == "min":

                        passed = (
                            value >= threshold
                        )

                    else:

                        passed = (
                            value <= threshold
                        )

                    ws.cell(
                        row=row,
                        column=excel_column
                    ).fill = (
                        green_fill
                        if passed
                        else red_fill
                    )

            # ------------------------------------------------
            # Freeze header
            # ------------------------------------------------

            ws.freeze_panes = "A2"

            # ------------------------------------------------
            # Autofilter
            # ------------------------------------------------

            ws.auto_filter.ref = (
                ws.dimensions
            )

            # ------------------------------------------------
            # Column widths
            # ------------------------------------------------

            for column_cells in ws.columns:

                max_length = 0

                column_letter = (
                    column_cells[0].column_letter
                )

                for cell in column_cells:

                    if cell.value is not None:

                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )

                ws.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    30
                )

        workbook.save(
            OUTPUT_FILE
        )

    # ========================================================
    # RUN DAY 17
    # ========================================================

    def run(self):

        print()
        print("=" * 70)
        print("SPRINT 3 - DAY 17")
        print("COMPOSITE SCORE + EXPORT")
        print("=" * 70)

        # ----------------------------------------------------
        # Load
        # ----------------------------------------------------

        df = self.load_data()

        print(
            f"✓ Loaded rows: {len(df)}"
        )

        # ----------------------------------------------------
        # Composite score
        # ----------------------------------------------------

        df = self.calculate_composite_score(
            df
        )

        # ----------------------------------------------------
        # Run all presets
        # ----------------------------------------------------

        all_results = {}

        for preset_name in PRESETS:

            print()
            print(
                f"Running: {preset_name}"
            )

            result = self.apply_preset(
                df,
                preset_name
            )

            result = result.sort_values(
                "composite_quality_score",
                ascending=False
            )

            all_results[
                preset_name
            ] = result

            print(
                f"✓ Results: {len(result)}"
            )

        # ----------------------------------------------------
        # Excel
        # ----------------------------------------------------

        self.generate_excel(
            all_results
        )

        # ----------------------------------------------------
        # Summary
        # ----------------------------------------------------

        print()
        print("=" * 70)
        print("DAY 17 COMPLETE")
        print("=" * 70)

        for name, result in all_results.items():

            print(
                f"{name:<25} "
                f"{len(result)} companies"
            )

        print()
        print(
            f"Output: {OUTPUT_FILE}"
        )

        return all_results


# ============================================================
# MAIN
# ============================================================

def main():

    engine = CompositeEngine()

    engine.run()


if __name__ == "__main__":

    main()