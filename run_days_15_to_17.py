# run_days_15_to_17.py

import os
import yaml
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
import numpy as np
import pandas as pd

# ---------------------------------------------------------
# 1. SETUP DIRECTORY STRUCTURE & CONFIG
# ---------------------------------------------------------
os.makedirs("config", exist_ok=True)
os.makedirs("src/screener", exist_ok=True)
os.makedirs("output", exist_ok=True)

config_content = {
    "metrics": {
        "roe": {"name": "ROE", "type": "min"},
        "de_ratio": {"name": "D/E", "type": "max"},
        "fcf": {"name": "FCF", "type": "min"},
        "rev_cagr_5yr": {"name": "Revenue CAGR 5yr", "type": "min"},
        "pat_cagr_5yr": {"name": "PAT CAGR 5yr", "type": "min"},
        "opm": {"name": "OPM", "type": "min"},
        "pe_ratio": {"name": "P/E", "type": "max"},
        "pb_ratio": {"name": "P/B", "type": "max"},
        "div_yield": {"name": "Dividend Yield", "type": "min"},
        "div_payout": {"name": "Dividend Payout", "type": "max"},
        "icr": {"name": "ICR", "type": "min"},
        "market_cap": {"name": "Market Cap", "type": "min"},
        "net_profit": {"name": "Net Profit", "type": "min"},
        "eps_cagr": {"name": "EPS CAGR", "type": "min"},
        "asset_turnover": {"name": "Asset Turnover", "type": "min"},
        "sales": {"name": "Sales", "type": "min"},
        "revenue": {"name": "Revenue", "type": "min"},
        "rev_cagr_3yr": {"name": "Revenue CAGR 3yr", "type": "min"},
        "fcf_latest": {"name": "Latest FCF", "type": "min"}
    },
    "sector_exceptions": {
        "de_ratio_skip_sectors": ["Financials", "Banking", "NBFC"]
    },
    "presets": {
        "quality_compounder": {
            "name": "Quality Compounder",
            "criteria": {"roe_min": 15.0, "de_ratio_max": 1.0, "fcf_min": 0.0, "rev_cagr_5yr_min": 10.0}
        },
        "value_pick": {
            "name": "Value Pick",
            "criteria": {"pe_ratio_max": 20.0, "pb_ratio_max": 3.0, "de_ratio_max": 2.0, "div_yield_min": 1.0}
        },
        "growth_accelerator": {
            "name": "Growth Accelerator",
            "criteria": {"pat_cagr_5yr_min": 20.0, "rev_cagr_5yr_min": 15.0, "de_ratio_max": 2.0}
        },
        "dividend_champion": {
            "name": "Dividend Champion",
            "criteria": {"div_yield_min": 2.0, "div_payout_max": 80.0, "fcf_min": 0.0}
        },
        "debt_free_blue_chip": {
            "name": "Debt-Free Blue Chip",
            "criteria": {"de_ratio_max": 0.0, "roe_min": 12.0, "revenue_min": 5000.0}
        },
        "turnaround_watch": {
            "name": "Turnaround Watch",
            "criteria": {"rev_cagr_3yr_min": 10.0, "fcf_latest_min": 0.0, "de_declining_yoy": True}
        }
    }
}

with open("config/screener_config.yaml", "w") as f:
    yaml.dump(config_content, f, default_flow_style=False)

print("[✓] Created config/screener_config.yaml")

# ---------------------------------------------------------
# 2. SCREENER ENGINE CLASS DEFINITION
# ---------------------------------------------------------
class ScreenerEngine:
    def __init__(self, config_path: str = "config/screener_config.yaml"):
        with open(config_path, "r") as file:
            self.config = yaml.safe_load(file)

    def preprocess_data(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        if 'icr' in df.columns:
            df['icr'] = df['icr'].replace(
                ['Debt Free', 'DEBT FREE', 'debt free', 'N/A', None], np.inf
            )
            df['icr'] = pd.to_numeric(df['icr'], errors='coerce').fillna(np.inf)
        return df

    def winsorize_and_scale(self, series: pd.Series, lower_p: float = 0.10, upper_p: float = 0.90, invert: bool = False) -> pd.Series:
        clean_series = series.replace([np.inf, -np.inf], np.nan)
        p10 = clean_series.quantile(lower_p)
        p90 = clean_series.quantile(upper_p)
        
        if p10 == p90:
            return pd.Series(50.0, index=series.index)
            
        clipped = clean_series.clip(lower=p10, upper=p90)
        if invert:
            scaled = 100.0 * (p90 - clipped) / (p90 - p10)
        else:
            scaled = 100.0 * (clipped - p10) / (p90 - p10)
            
        return scaled.fillna(0.0)

    def calculate_composite_score(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        sector_col = 'broad_sector' if 'broad_sector' in df.columns else 'sector'
        
        df['s_roe'] = 0.0
        df['s_roce'] = 0.0
        df['s_npm'] = 0.0
        df['s_fcf_cagr'] = 0.0
        df['s_cfo_pat'] = 0.0
        df['s_fcf_flag'] = np.where(df.get('fcf', 0) > 0, 100.0, 0.0)
        df['s_rev_cagr'] = 0.0
        df['s_pat_cagr'] = 0.0
        df['s_de'] = 0.0
        df['s_icr'] = 0.0

        for sector_name, group in df.groupby(sector_col):
            idx = group.index
            if 'roe' in df.columns:
                df.loc[idx, 's_roe'] = self.winsorize_and_scale(group['roe'])
            if 'roce' in df.columns:
                df.loc[idx, 's_roce'] = self.winsorize_and_scale(group['roce'])
            if 'npm' in df.columns:
                df.loc[idx, 's_npm'] = self.winsorize_and_scale(group['npm'])
            if 'fcf_cagr_5yr' in df.columns:
                df.loc[idx, 's_fcf_cagr'] = self.winsorize_and_scale(group['fcf_cagr_5yr'])
            if 'cfo_pat_ratio' in df.columns:
                df.loc[idx, 's_cfo_pat'] = self.winsorize_and_scale(group['cfo_pat_ratio'])
            if 'rev_cagr_5yr' in df.columns:
                df.loc[idx, 's_rev_cagr'] = self.winsorize_and_scale(group['rev_cagr_5yr'])
            if 'pat_cagr_5yr' in df.columns:
                df.loc[idx, 's_pat_cagr'] = self.winsorize_and_scale(group['pat_cagr_5yr'])
            if 'de_ratio' in df.columns:
                df.loc[idx, 's_de'] = self.winsorize_and_scale(group['de_ratio'], invert=True)
            if 'icr' in df.columns:
                df.loc[idx, 's_icr'] = self.winsorize_and_scale(group['icr'])

        df['profitability_score'] = (0.15 * df['s_roe']) + (0.10 * df['s_roce']) + (0.10 * df['s_npm'])
        df['cash_quality_score']  = (0.15 * df['s_fcf_cagr']) + (0.10 * df['s_cfo_pat']) + (0.05 * df['s_fcf_flag'])
        df['growth_score']        = (0.10 * df['s_rev_cagr']) + (0.10 * df['s_pat_cagr'])
        df['leverage_score']      = (0.10 * df['s_de']) + (0.05 * df['s_icr'])

        df['composite_quality_score'] = (
            df['profitability_score'] + df['cash_quality_score'] + 
            df['growth_score'] + df['leverage_score']
        ).round(2)

        temp_cols = ['s_roe', 's_roce', 's_npm', 's_fcf_cagr', 's_cfo_pat', 
                     's_fcf_flag', 's_rev_cagr', 's_pat_cagr', 's_de', 's_icr',
                     'profitability_score', 'cash_quality_score', 'growth_score', 'leverage_score']
        df.drop(columns=temp_cols, inplace=True, errors='ignore')
        return df

    def apply_filters(self, df: pd.DataFrame, criteria: dict) -> pd.DataFrame:
        df = self.preprocess_data(df)
        df = self.calculate_composite_score(df)
        
        skip_de_sectors = self.config.get('sector_exceptions', {}).get('de_ratio_skip_sectors', [])
        mask = pd.Series(True, index=df.index)
        
        for key, threshold in criteria.items():
            if threshold is None:
                continue
            if key == 'de_declining_yoy' and threshold is True:
                if 'de_ratio_yoy_change' in df.columns:
                    mask &= (df['de_ratio_yoy_change'] < 0)
                continue

            if key.endswith('_min'):
                metric = key[:-4]
                if metric in df.columns:
                    mask &= (df[metric] >= threshold)
            elif key.endswith('_max'):
                metric = key[:-4]
                if metric in df.columns:
                    if metric == 'de_ratio':
                        sector_col = 'broad_sector' if 'broad_sector' in df.columns else 'sector'
                        if sector_col in df.columns:
                            is_financial = df[sector_col].isin(skip_de_sectors)
                            passes_de = (df['de_ratio'] <= threshold) | is_financial
                            mask &= passes_de
                        else:
                            mask &= (df['de_ratio'] <= threshold)
                    else:
                        mask &= (df[metric] <= threshold)

        filtered_df = df[mask].copy()
        return filtered_df.sort_values(by='composite_quality_score', ascending=False)

    def run_preset(self, df: pd.DataFrame, preset_key: str) -> pd.DataFrame:
        presets = self.config.get('presets', {})
        if preset_key not in presets:
            raise ValueError(f"Preset '{preset_key}' not found in configuration.")
        return self.apply_filters(df, presets[preset_key]['criteria'])

# ---------------------------------------------------------
# 3. GENERATE UNIVERSE DATA & RUN PIPELINE
# ---------------------------------------------------------
np.random.seed(42)
sectors = ["IT Services", "Financials", "Consumer Goods", "Manufacturing", "Pharma"]

mock_universe = pd.DataFrame({
    "company_id": [f"COMP_{i:02d}" for i in range(1, 93)],
    "company_name": [f"Company {i}" for i in range(1, 93)],
    "broad_sector": np.random.choice(sectors, 92),
    "roe": np.random.uniform(5, 30, 92),
    "roce": np.random.uniform(5, 25, 92),
    "npm": np.random.uniform(2, 20, 92),
    "de_ratio": np.random.uniform(0, 3, 92),
    "icr": np.random.choice([2.5, 5.0, 10.0, "Debt Free"], 92),
    "fcf": np.random.uniform(-50, 500, 92),
    "fcf_latest": np.random.uniform(-10, 200, 92),
    "fcf_cagr_5yr": np.random.uniform(-5, 25, 92),
    "cfo_pat_ratio": np.random.uniform(0.5, 1.5, 92),
    "rev_cagr_5yr": np.random.uniform(2, 25, 92),
    "rev_cagr_3yr": np.random.uniform(2, 20, 92),
    "pat_cagr_5yr": np.random.uniform(2, 30, 92),
    "pe_ratio": np.random.uniform(8, 45, 92),
    "pb_ratio": np.random.uniform(0.8, 6.0, 92),
    "div_yield": np.random.uniform(0, 4.0, 92),
    "div_payout": np.random.uniform(10, 90, 92),
    "revenue": np.random.uniform(1000, 20000, 92),
    "de_ratio_yoy_change": np.random.uniform(-0.5, 0.5, 92)
})

engine = ScreenerEngine()

print("\n--- Day 16 Preset Validation ---")
presets = engine.config.get('presets', {})
for key, info in presets.items():
    res = engine.run_preset(mock_universe, key)
    count = len(res)
    status = "PASSED" if 5 <= count <= 50 else "FAILED"
    print(f"[{status}] {info['name']}: {count} companies returned")

# ---------------------------------------------------------
# 4. EXPORT TO screener_output.xlsx
# ---------------------------------------------------------
output_path = "output/screener_output.xlsx"
wb = openpyxl.Workbook()
wb.remove(wb.active)

GREEN_FILL = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
RED_FILL   = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

KPI_COLUMNS = [
    'company_id', 'company_name', 'broad_sector', 'composite_quality_score',
    'roe', 'roce', 'npm', 'de_ratio', 'icr', 'fcf', 'fcf_cagr_5yr',
    'cfo_pat_ratio', 'rev_cagr_5yr', 'rev_cagr_3yr', 'pat_cagr_5yr',
    'pe_ratio', 'pb_ratio', 'div_yield', 'div_payout', 'revenue'
]

for preset_key, preset_info in presets.items():
    preset_name = preset_info['name']
    criteria = preset_info['criteria']
    filtered_df = engine.apply_filters(mock_universe, criteria)
    cols_to_write = [c for c in KPI_COLUMNS if c in filtered_df.columns]
    
    ws = wb.create_sheet(title=preset_name[:31])
    ws.views.sheetView[0].showGridLines = True
    
    ws.append(cols_to_write)
    for col_idx in range(1, len(cols_to_write) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_idx, record in enumerate(filtered_df[cols_to_write].to_dict('records'), start=2):
        for col_idx, col_name in enumerate(cols_to_write, start=1):
            val = record[col_name]
            cell = ws.cell(row=row_idx, column=col_idx, value=val if val != float('inf') else "Debt Free")
            
            min_key = f"{col_name}_min"
            max_key = f"{col_name}_max"
            
            if min_key in criteria and isinstance(val, (int, float)):
                cell.fill = GREEN_FILL if val >= criteria[min_key] else RED_FILL
            elif max_key in criteria and isinstance(val, (int, float)):
                cell.fill = GREEN_FILL if val <= criteria[max_key] else RED_FILL
            elif col_name == 'de_ratio' and 'de_ratio_max' in criteria:
                if record.get('broad_sector') in ["Financials", "Banking", "NBFC"]:
                    cell.fill = GREEN_FILL
                elif isinstance(val, (int, float)):
                    cell.fill = GREEN_FILL if val <= criteria['de_ratio_max'] else RED_FILL

wb.save(output_path)
print(f"\n[✓] Generated {output_path} with 6 color-coded preset sheets.")